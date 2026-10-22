"""
Monthly Business-Finance MIS pipeline.

    raw CSVs ──► data-quality checks ──► SQLite warehouse ──► SQL KPI layer
             ──► Claude commentary (number-verified) ──► Excel report + Markdown summary + charts

Run:  python -m src.pipeline --month 2026-09
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import time
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.ai_commentary import build_facts, generate_commentary
from src.quality_checks import validate

ROOT = Path(__file__).resolve().parents[1]
RAW, OUT = ROOT / "data" / "raw", ROOT / "reports"
SQL_FILE = ROOT / "sql" / "kpi_queries.sql"
DB = ROOT / "data" / "warehouse.db"


def load_queries() -> dict[str, str]:
    text = SQL_FILE.read_text(encoding="utf-8")
    parts = re.split(r"^-- name: (\w+)\s*$", text, flags=re.M)
    return {parts[i]: parts[i + 1].strip() for i in range(1, len(parts), 2)}


def build_warehouse(month: str) -> pd.DataFrame:
    tx = pd.read_csv(RAW / "transactions.csv")
    merchants = pd.read_csv(RAW / "merchants.csv")
    refunds = pd.read_csv(RAW / "refunds.csv")
    as_of = pd.Timestamp(month + "-01") + pd.DateOffset(months=1) - pd.Timedelta(seconds=1)

    clean_tx, clean_ref, quarantine, dq_log = validate(tx, merchants, refunds, as_of)

    with sqlite3.connect(DB) as con:
        clean_tx.assign(created_at=clean_tx.created_at.dt.strftime("%Y-%m-%d %H:%M:%S")) \
            .to_sql("transactions", con, if_exists="replace", index=False)
        merchants.to_sql("merchants", con, if_exists="replace", index=False)
        clean_ref.assign(refunded_at=clean_ref.refunded_at.dt.strftime("%Y-%m-%d %H:%M:%S")) \
            .to_sql("refunds", con, if_exists="replace", index=False)
        quarantine.astype({"created_at": str}).to_sql("dq_quarantine", con, if_exists="replace", index=False)
        dq_log.to_sql("dq_log", con, if_exists="replace", index=False)
        con.executescript("""
            CREATE INDEX IF NOT EXISTS ix_tx_created ON transactions(created_at);
            CREATE INDEX IF NOT EXISTS ix_tx_merchant ON transactions(merchant_id);
        """)
    print(f"[dq] input={len(tx):,} clean={len(clean_tx):,} quarantined={len(quarantine):,}")
    return dq_log


def run_kpis(month: str) -> dict[str, pd.DataFrame]:
    queries = load_queries()
    with sqlite3.connect(DB) as con:
        return {name: pd.read_sql_query(q, con, params={"report_month": month}) for name, q in queries.items()}


# ---------------------------------------------------------------- charts
def make_charts(k: dict[str, pd.DataFrame], month: str) -> list[Path]:
    OUT.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"svg.fonttype": "none", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
    paths = []

    pnl = k["monthly_pnl"]
    fig, ax1 = plt.subplots(figsize=(9, 4.2))
    ax1.bar(pnl.month, pnl.gmv / 1e7, color="#3B6FB6", label="GMV (₹ Cr)")
    ax1.set_ylabel("GMV (₹ Cr)")
    ax1.set_ylim(0, pnl.gmv.max() / 1e7 * 1.3)
    ax2 = ax1.twinx()
    ax2.plot(pnl.month, pnl.take_rate_bps, color="#D9822B", marker="o", lw=2, label="Take rate (bps)")
    ax2.set_ylabel("Take rate (bps)")
    ax2.spines["top"].set_visible(False)
    ax1.set_title("GMV and take rate by month")
    fig.legend(loc="upper left", bbox_to_anchor=(0.08, 0.88), frameon=False)
    fig.tight_layout()
    p = OUT / f"gmv_take_rate_{month}.svg"
    fig.savefig(p); plt.close(fig); paths.append(p)

    mm = k["method_mix"].sort_values("contribution")
    fig, ax = plt.subplots(figsize=(8, 3.8))
    colors = ["#C0392B" if v < 0 else "#2E8B57" for v in mm.contribution]
    ax.barh(mm.method, mm.contribution / 1e5, color=colors)
    ax.axvline(0, color="#555", lw=0.8)
    ax.set_xlabel("Contribution (₹ Lakh)")
    ax.set_title(f"Contribution by payment method — {month}")
    fig.tight_layout()
    p = OUT / f"method_contribution_{month}.svg"
    fig.savefig(p); plt.close(fig); paths.append(p)
    return paths


# ---------------------------------------------------------------- excel
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INR_FMT = '[>=10000000]"₹"#,##,##,##0.00;[>=100000]"₹"#,##,##0.00;"₹"#,##0.00'


def _style_sheet(ws, money_cols: set[str]) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for col_idx, cell in enumerate(ws[1], start=1):
        letter = get_column_letter(col_idx)
        width = max(len(str(cell.value)), *(len(str(c.value or "")) for c in ws[letter][1:50])) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 45)
        if cell.value in money_cols:
            for c in ws[letter][1:]:
                c.number_format = INR_FMT


def write_excel(k, dq_log, commentary, source, unverified, month) -> Path:
    path = OUT / f"finance_mis_{month}.xlsx"
    money = {"gmv", "refunds", "net_gmv", "gross_revenue", "network_cost", "contribution",
             "avg_ticket", "prev_month_gmv", "report_month_gmv"}
    sheets = {
        "Monthly P&L": k["monthly_pnl"],
        "Payment Methods": k["method_mix"],
        "Segments": k["segment_summary"],
        "Top Merchants": k["top_merchants"],
        "At-Risk Merchants": k["declining_merchants"],
        "Category Trend": k["category_trend"].pivot(index="month", columns="category", values="gmv").reset_index(),
        "Merchant Retention": k["merchant_retention"].pivot(index="cohort", columns="month_n",
                                                            values="retention_pct").reset_index(),
        "DQ Log": dq_log,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary = pd.DataFrame({"Monthly finance MIS": [
            f"Report month: {month}", f"Commentary source: {source}",
            "Unverified numbers in commentary: " + (", ".join(unverified) if unverified else "none — all figures traced to source"),
            "", *commentary.replace("**", "").splitlines()]})
        summary.to_excel(xw, sheet_name="Summary", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)

    wb = load_workbook(path)
    ws = wb["Summary"]
    ws.column_dimensions["A"].width = 120
    ws["A1"].fill, ws["A1"].font = HEADER_FILL, HEADER_FONT
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in sheets:
        _style_sheet(wb[name], money)

    # native Excel charts so the workbook stays editable for finance
    ws = wb["Monthly P&L"]
    n = ws.max_row
    bar = BarChart(); bar.title = "GMV by month"; bar.y_axis.title = "INR"
    bar.add_data(Reference(ws, min_col=4, min_row=1, max_row=n), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    bar.height, bar.width = 8, 16
    ws.add_chart(bar, f"A{n + 3}")
    line = LineChart(); line.title = "Take rate (bps)"
    line.add_data(Reference(ws, min_col=11, min_row=1, max_row=n), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    line.height, line.width = 8, 16
    ws.add_chart(line, f"J{n + 3}")

    ws = wb["DQ Log"]
    colors = {"PASS": "C6EFCE", "WARN": "FFEB9C", "FAIL": "FFC7CE"}
    status_col = [c.value for c in ws[1]].index("status") + 1
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, status_col)
        c.fill = PatternFill("solid", fgColor=colors.get(c.value, "FFFFFF"))
    wb.save(path)
    return path


def write_markdown(k, dq_log, commentary, source, unverified, month, charts) -> Path:
    pnl = k["monthly_pnl"]
    cur = pnl.iloc[-1]
    path = OUT / f"finance_mis_{month}.md"
    lines = [
        f"# Business Finance MIS — {month}", "",
        f"*Commentary source: {source}*  ",
        f"*Unverified numbers: {', '.join(unverified) if unverified else 'none — every figure traced to the KPI layer'}*", "",
        "## Commentary", "", commentary, "",
        "## Key metrics", "",
        "| Metric | Value |", "|---|---|",
        f"| GMV | ₹{cur.gmv/1e7:,.2f} Cr ({cur.gmv_mom_pct:+}% MoM) |",
        f"| Gross revenue | ₹{cur.gross_revenue/1e5:,.2f} L ({cur.revenue_mom_pct:+}% MoM) |",
        f"| Contribution | ₹{cur.contribution/1e5:,.2f} L ({cur.contribution_margin_pct}% margin) |",
        f"| Take rate | {cur.take_rate_bps} bps |",
        f"| Success rate | {cur.success_rate_pct}% |",
        f"| Active merchants | {int(cur.active_merchants):,} |", "",
        *[f"![{c.stem}]({c.name})" for c in charts], "",
        "## Data-quality log", "", dq_log.to_markdown(index=False),
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main(month: str) -> None:
    t0 = time.time()
    OUT.mkdir(parents=True, exist_ok=True)
    dq_log = build_warehouse(month)
    if (dq_log.check.str.startswith("Control total") & dq_log.status.eq("FAIL")).any():
        raise SystemExit("Control totals do not tie out — report NOT published.")
    k = run_kpis(month)
    facts = build_facts(k["monthly_pnl"], k["method_mix"], k["segment_summary"],
                        k["declining_merchants"], dq_log)
    commentary, source, unverified = generate_commentary(facts)
    charts = make_charts(k, month)
    xlsx = write_excel(k, dq_log, commentary, source, unverified, month)
    md = write_markdown(k, dq_log, commentary, source, unverified, month, charts)
    print(f"[done] {xlsx.name}, {md.name} + {len(charts)} charts in {time.time() - t0:.1f}s")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", default="2026-09", help="Report month YYYY-MM")
    main(ap.parse_args().month)
